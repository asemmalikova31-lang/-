
import http.server, socketserver, json, os, re, hashlib, datetime, urllib.parse, mimetypes, io, zipfile, xml.etree.ElementTree as ET
from pathlib import Path

ROOT=Path(__file__).resolve().parent.parent
CASE_DIR=ROOT/"cases"; LAW_DIR=ROOT/"law"
PORT=8765
try:
    from pypdf import PdfReader
except Exception: PdfReader=None
try:
    from docx import Document
except Exception: Document=None
try:
    import openpyxl
except Exception: openpyxl=None

def sha256(b): return hashlib.sha256(b).hexdigest()
def norm(s): return re.sub(r'\s+',' ',s or '').strip()
DATE_PAT=re.compile(r'(?<!\d)(\d{1,2})[./-](\d{1,2})[./-](20\d{2})(?!\d)')
def dates(text):
    out=[]
    for m in DATE_PAT.finditer(text or ''):
        try: out.append(datetime.date(int(m.group(3)),int(m.group(2)),int(m.group(1))))
        except: pass
    return out

def extract(path):
    ext=path.suffix.lower(); text=''
    if ext=='.pdf' and PdfReader:
        try:
            r=PdfReader(str(path))
            text='\n'.join((p.extract_text() or '') for p in r.pages)
        except Exception as e: text=f'[PDF_ERROR] {e}'
    elif ext=='.docx' and Document:
        try:
            d=Document(str(path))
            text='\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
        except Exception as e: text=f'[DOCX_ERROR] {e}'
    elif ext=='.xlsx' and openpyxl:
        try:
            wb=openpyxl.load_workbook(path,data_only=True,read_only=True)
            text='\n'.join(' | '.join('' if v is None else str(v) for v in row) for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
        except Exception as e: text=f'[XLSX_ERROR] {e}'
    elif ext in ('.txt','.md','.csv','.json'):
        text=path.read_text(encoding='utf-8',errors='ignore')
    else:
        # DOC/XLS/RTF/images are recorded for the case; legacy binary/image OCR is optional.
        text=''
    return norm(text)

def classify(name,text):
    x=(name+' '+text).lower()
    rules=[
      ('Приказ',r'приказ|распоряжение|наложить.{0,30}взыскани|привлечь.{0,30}дисциплин'),
      ('Объяснительная',r'объяснитель|пояснитель'),
      ('Требование объяснения',r'требован.{0,30}объясн|предостав.{0,30}объясн|запрос.{0,30}объясн'),
      ('Акт проверки',r'акт.{0,30}(провер|служебн|расслед)|акт проверки'),
      ('Заключение проверки',r'заключен.{0,30}(провер|расслед|служебн)'),
      ('Обращение',r'обращен|жалоб|заявлен.{0,30}(граждан|физичес)'),
      ('СЗ/докладная',r'служебн.{0,30}записк|докладн|рапорт'),
      ('Проверка/расследование',r'провер|служебн.{0,30}расслед'),
    ]
    for k,p in rules:
        if re.search(p,x): return k
    return 'Иной документ'

def substantive(text):
    return bool(re.search(r'\b(работник|сотрудник|проступ|наруш|не испол|не выпол|установлен|выявлен|винов|дисциплин)\b', text, re.I))

def primary_reason(doc):
    return ("Документ содержит признаки конкретного нарушения/проступка и "
            "идентифицируемого работника; он рассматривается как кандидат на первичный "
            "источник. Дата из приказа или производной СЗ автоматически не принимается.")

def analyze(docs):
    # Candidate source: substantive source documents, excluding purely procedural docs.
    c=[d for d in docs if d["kind"] in ["Акт проверки","Заключение проверки","Обращение","СЗ/докладная","Проверка/расследование"] and substantive(d["text"])]
    c.sort(key=lambda d: (d["date"] or "9999-99-99", d["name"]))
    primary=c[0] if c else None
    order=next((d for d in docs if d["kind"]=="Приказ"),None)
    req=next((d for d in docs if d["kind"]=="Требование объяснения"),None)
    expl=next((d for d in docs if d["kind"]=="Объяснительная"),None)
    def add30(s):
        if not s:return None
        y,m,d=map(int,s.split('-')); return (datetime.date(y,m,d)+datetime.timedelta(days=30)).isoformat()
    deadline=add30(primary["date"] if primary else None)
    findings=[]
    findings.append({"status":"ok" if primary else "bad","title":"Первичный источник обнаружения",
                     "text":(primary["name"]+"; "+primary["date"] if primary else "Не найден"),
                     "reason":primary_reason(primary) if primary else "Нужен первичный материал."})
    if primary and order:
        findings.append({"status":"ok" if order["date"]<=deadline else "bad","title":"Месячный срок",
                         "text":f"Обнаружение: {primary['date']}; ориентир +30 дней: {deadline}; приказ: {order['date']}",
                         "reason":"Расчёт является предварительным; юрист проверяет применимое правило и исключения."})
    else: findings.append({"status":"warn","title":"Месячный срок","text":"Недостаточно подтверждённых дат","reason":"Требуется ручная проверка."})
    findings += [
      {"status":"ok" if req else "bad","title":"Требование объяснения","text":req["name"] if req else "Не найдено","reason":"Проверяется ст. 65 ТК РК и подтверждение получения."},
      {"status":"ok" if expl else "warn","title":"Объяснение работника","text":expl["name"] if expl else "Не найдено","reason":"Если объяснение отсутствует, требуется проверить акт/подтверждение непредставления."},
    ]
    return {"primary":primary,"order":order,"deadline":deadline,"findings":findings}

def load_law():
    arr=[]
    for p in LAW_DIR.glob('*.json'):
        try: arr.append(json.loads(p.read_text(encoding='utf-8')))
        except: pass
    return arr

class Handler(http.server.BaseHTTPRequestHandler):
    def send(self,code,ctype,b):
        self.send_response(code); self.send_header('Content-Type',ctype); self.send_header('Content-Length',str(len(b))); self.end_headers(); self.wfile.write(b)
    def do_GET(self):
        u=urllib.parse.urlparse(self.path)
        if u.path=='/api/law':
            self.send(200,'application/json; charset=utf-8',json.dumps(load_law(),ensure_ascii=False).encode()); return
        if u.path=='/api/cases':
            items=[]
            for p in CASE_DIR.glob('*.json'):
                try: items.append(json.loads(p.read_text(encoding='utf-8')))
                except: pass
            self.send(200,'application/json; charset=utf-8',json.dumps(items,ensure_ascii=False).encode()); return
        if u.path=='/api/file':
            q=urllib.parse.parse_qs(u.query); p=Path(q.get('path',[''])[0])
            if p.is_file() and ROOT in p.resolve().parents:
                self.send(200,'application/octet-stream',p.read_bytes()); return
            self.send(404,'text/plain',b'Not found'); return
        p=ROOT/'legalcheck'/'index.html'
        self.send(200,'text/html; charset=utf-8',p.read_bytes())
    def do_POST(self):
        u=urllib.parse.urlparse(self.path); n=int(self.headers.get('Content-Length','0')); body=self.rfile.read(n)
        if u.path=='/api/analyze':
            data=json.loads(body); docs=data.get('documents',[])
            result=analyze(docs)
            self.send(200,'application/json; charset=utf-8',json.dumps(result,ensure_ascii=False).encode()); return
        if u.path=='/api/law':
            data=json.loads(body); name=re.sub(r'[^a-zA-Z0-9_.-]','_',data.get('filename','law.json'))
            (LAW_DIR/name).write_text(json.dumps(data.get('content',{}),ensure_ascii=False,indent=2),encoding='utf-8')
            self.send(200,'application/json',b'{"ok":true}'); return
        if u.path=='/api/case':
            data=json.loads(body); cid=data.get('case_id') or ('LCZ-'+datetime.datetime.now().strftime('%Y%m%d%H%M%S'))
            (CASE_DIR/(cid+'.json')).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
            self.send(200,'application/json',json.dumps({"ok":True,"case_id":cid}).encode()); return
        self.send(404,'text/plain',b'Not found')

if __name__=='__main__':
    os.chdir(ROOT)
    with socketserver.ThreadingTCPServer(('127.0.0.1',PORT),Handler) as s:
        print(f'LegalCheck KZ: http://127.0.0.1:{PORT}')
        s.serve_forever()
